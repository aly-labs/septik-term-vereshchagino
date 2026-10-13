#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Локальный сервер сайта «ТЕРМИТ Верещагино».

Что делает:
  1) Раздаёт статические файлы (index.html, css, js, data, img).
  2) Принимает заявки на POST /api/lead.
     - Всегда сохраняет заявку в leads.jsonl (чтобы она не потерялась).
     - Дописывает её в Excel-таблицу «Заявки.xlsx».
     - Если задан YADISK_TOKEN — загружает таблицу на Яндекс.Диск
       (файл «Заявки с сайта.xlsx» в корне диска, открывается как таблица).
     - Если задан N8N_WEBHOOK_URL — дублирует заявку в n8n (Telegram/CRM).
  3) ЧПУ: /catalog/<slug> отдаёт index.html и подставляет в <title>/<meta>
     название модели для поисковиков.

СЕКРЕТЫ ХРАНЯТСЯ ТОЛЬКО ЗДЕСЬ, В ПЕРЕМЕННЫХ ОКРУЖЕНИЯ, И НЕ ПОПАДАЮТ В БРАУЗЕР.

Запуск:
    python server.py
    открыть http://localhost:8000

Переменные окружения:
    PORT               — порт (по умолчанию 8000)
    YADISK_TOKEN       — OAuth-токен Яндекс.Диска (как получить — см. yadisk_token.py)
    N8N_WEBHOOK_URL    — URL вебхука n8n (если пусто — только локальное сохранение)
    N8N_AUTH_HEADER    — необязательный заголовок Authorization для вебхука

Пример для Windows PowerShell:
    $env:YADISK_TOKEN = "ваш-токен"
    $env:N8N_WEBHOOK_URL = "https://ваш-домен/webhook/ваш-id"
    python server.py
"""

import csv
import datetime
import json
import os
import re
import ssl
import sys
import urllib.error
import urllib.parse
import urllib.request
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PORT = int(os.environ.get("PORT", "8000"))
WEBHOOK_URL = (os.environ.get("N8N_WEBHOOK_URL") or "").strip()
AUTH_HEADER = (os.environ.get("N8N_AUTH_HEADER") or "").strip()

# --- Яндекс.Диск ---
YADISK_TOKEN = (os.environ.get("YADISK_TOKEN") or "").strip()
LEADS_XLSX = os.path.join(BASE_DIR, "Заявки.xlsx")
LEADS_CSV = os.path.join(BASE_DIR, "leads.csv")
YADISK_REMOTE = "Заявки с сайта.xlsx"  # имя файла на Яндексе

LEADS_FILE = os.path.join(BASE_DIR, "leads.jsonl")
CATALOG_FILE = os.path.join(BASE_DIR, "data", "catalog.json")

TITLE_RE = re.compile(r"<title>.*?</title>", re.S)
DESC_RE = re.compile(r'<meta name="description" content=".*?"', re.S)

LEAD_HEADERS = [
    "Дата и время", "Тип заявки", "Модель / объект", "Имя", "Телефон",
    "Населённый пункт", "Комментарий", "Ответы калькулятора", "Страница",
]

TYPE_NAMES = {
    "callback": "Заказ звонка",
    "selection": "Подбор септика",
    "montage": "Расчёт монтажа",
    "consult": "Консультация",
    "delivery": "Расчёт доставки",
    "model": "Стоимость с монтажом",
    "calculator": "Калькулятор",
    "test": "Тест",
}


def load_catalog():
    try:
        with open(CATALOG_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as exc:  # noqa: BLE001
        print("[catalog] не удалось прочитать catalog.json:", exc, file=sys.stderr)
        return {"items": []}


CATALOG = load_catalog()


def log(*args):
    print("[server]", *args, flush=True)


def lead_row(lead):
    answers = lead.get("answers") or {}
    if isinstance(answers, dict) and answers:
        answers_str = "; ".join(f"{k}: {v}" for k, v in answers.items())
    else:
        answers_str = ""
    return [
        str(lead.get("receivedAt", "")).replace("T", " "),
        TYPE_NAMES.get(lead.get("type", ""), lead.get("type", "")),
        lead.get("model", ""),
        lead.get("name", ""),
        lead.get("phone", ""),
        lead.get("city", ""),
        lead.get("comment", ""),
        answers_str,
        lead.get("page", ""),
    ]


def append_to_table(lead):
    """Дописывает заявку в Excel-таблицу (или CSV, если openpyxl недоступен)."""
    row = lead_row(lead)
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        return append_csv(row)
    try:
        if os.path.exists(LEADS_XLSX):
            wb = load_workbook(LEADS_XLSX)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Заявки"
            ws.append(LEAD_HEADERS)
            for col, width in enumerate([20, 18, 24, 16, 18, 20, 32, 44, 30], start=1):
                ws.column_dimensions[get_column_letter(col)].width = width
            ws.freeze_panes = "A2"
        ws.append(row)
        wb.save(LEADS_XLSX)
        return "xlsx"
    except Exception as exc:  # noqa: BLE001
        log("ОШИБКА записи Excel:", exc)
        return append_csv(row)


def append_csv(row):
    try:
        new_file = not os.path.exists(LEADS_CSV)
        with open(LEADS_CSV, "a", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            if new_file:
                w.writerow(LEAD_HEADERS)
            w.writerow(row)
        return "csv"
    except OSError as exc:
        log("ОШИБКА записи CSV:", exc)
        return None


def upload_to_yadisk(local_path):
    """Загружает таблицу на Яндекс.Диск WebDAV. Возвращает True/False/None."""
    if not YADISK_TOKEN:
        log("YADISK_TOKEN не задан — таблица только локально:", local_path)
        return None
    try:
        url = "https://webdav.yandex.ru/" + urllib.parse.quote(YADISK_REMOTE)
        with open(local_path, "rb") as f:
            data = f.read()
        req = urllib.request.Request(
            url, data=data, method="PUT",
            headers={
                "Authorization": "OAuth " + YADISK_TOKEN,
                "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            },
        )
        with urllib.request.urlopen(req, timeout=30, context=ssl.create_default_context()) as resp:
            ok = 200 <= resp.status < 300
        log("Яндекс.Диск:", "таблица загружена" if ok else f"статус {resp.status}")
        return ok
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", "ignore")[:200]
        log("ОШИБКА Яндекс.Диска:", exc.code, exc.reason, body)
        return False
    except Exception as exc:  # noqa: BLE001
        log("ОШИБКА Яндекс.Диска:", exc)
        return False


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    # ---------- утилиты ----------
    def send_json(self, code, obj):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def send_index(self, slug=None, inject_base=False):
        try:
            with open(os.path.join(BASE_DIR, "index.html"), "r", encoding="utf-8") as f:
                html = f.read()
        except OSError:
            self.send_error(404, "index.html не найден")
            return

        if inject_base:
            html = html.replace("<head>", '<head>\n  <base href="/">', 1)

        if slug:
            item = next((i for i in CATALOG.get("items", []) if i.get("slug") == slug), None)
            if item:
                title = f'{item["name"]} — цена и монтаж в Верещагино | ТЕРМИТ'
                desc = (f'{item["name"]}: характеристики, цена на оборудование и расчёт монтажа '
                        f'в Верещагино и районе. Оставьте заявку на консультацию.')
                html = TITLE_RE.sub(f"<title>{title}</title>", html, count=1)
                html = DESC_RE.sub(f'<meta name="description" content="{desc}"', html, count=1)

        body = html.encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    # ---------- GET ----------
    def do_GET(self):
        path = self.path.split("?", 1)[0]

        if path in ("/", "/index.html"):
            self.send_index()
            return

        # ЧПУ страниц каталога: /catalog/<slug>
        m = re.match(r"^/catalog/([^/]+)/?$", path)
        if m:
            self.send_index(slug=urllib.parse.unquote(m.group(1)), inject_base=True)
            return

        # SPA-фолбэк: неизвестный путь без расширения → index.html
        if "." not in os.path.basename(path):
            self.send_index(inject_base=True)
            return

        return super().do_GET()

    # ---------- POST ----------
    def do_POST(self):
        path = self.path.split("?", 1)[0]
        if path != "/api/lead":
            self.send_json(404, {"ok": False, "error": "Не найдено"})
            return

        try:
            length = int(self.headers.get("Content-Length", "0"))
            raw = self.rfile.read(length).decode("utf-8") if length else ""
            data = json.loads(raw) if raw else {}
            if not isinstance(data, dict):
                raise ValueError("ожидался объект JSON")
        except Exception as exc:  # noqa: BLE001
            log("отклонена некорректная заявка:", exc)
            self.send_json(400, {"ok": False, "error": "Некорректные данные"})
            return

        name = str(data.get("name", "")).strip()
        phone = str(data.get("phone", "")).strip()
        digits = re.sub(r"\D", "", phone)
        if not name or len(digits) < 10:
            self.send_json(422, {"ok": False, "error": "Укажите имя и корректный телефон"})
            return

        lead = {
            "receivedAt": datetime.datetime.now().isoformat(timespec="seconds"),
            "type": data.get("type", ""),
            "model": data.get("model", ""),
            "name": name,
            "phone": phone,
            "city": data.get("city", ""),
            "comment": data.get("comment", ""),
            "answers": data.get("answers", {}),
            "page": data.get("page", ""),
            "userAgent": self.headers.get("User-Agent", ""),
        }

        # 1) Всегда сохраняем локально — заявка не теряется.
        try:
            with open(LEADS_FILE, "a", encoding="utf-8") as f:
                f.write(json.dumps(lead, ensure_ascii=False) + "\n")
        except OSError as exc:
            log("ОШИБКА записи leads.jsonl:", exc)

        # 2) Дописываем в Excel/CSV и загружаем на Яндекс.Диск.
        fmt = None
        try:
            fmt = append_to_table(lead)
            upload_to_yadisk(LEADS_XLSX if os.path.exists(LEADS_XLSX) else LEADS_CSV)
        except Exception as exc:  # noqa: BLE001
            log("ОШИБКА таблицы заявок:", exc)

        # 3) Пересылаем в n8n, если настроено.
        forwarded = False
        if WEBHOOK_URL:
            try:
                req = urllib.request.Request(
                    WEBHOOK_URL,
                    data=json.dumps(lead, ensure_ascii=False).encode("utf-8"),
                    headers={"Content-Type": "application/json"},
                    method="POST",
                )
                if AUTH_HEADER:
                    req.add_header("Authorization", AUTH_HEADER)
                with urllib.request.urlopen(req, timeout=12, context=ssl.create_default_context()) as resp:
                    forwarded = 200 <= resp.status < 300
                if not forwarded:
                    log("n8n вернул статус", resp.status)
            except urllib.error.URLError as exc:
                log("ОШИБКА отправки в n8n:", exc)
            except Exception as exc:  # noqa: BLE001
                log("ОШИБКА отправки в n8n:", exc)
        else:
            log("N8N_WEBHOOK_URL не задан — заявка в leads.jsonl / таблице")

        log(f"заявка: {lead['type'] or 'заявка'} · {name} · {phone} · таблица={fmt} · n8n={'ok' if forwarded else 'нет'}")
        self.send_json(200, {"ok": True, "forwarded": forwarded, "table": bool(fmt)})

    def log_message(self, fmt, *args):  # тише в консоли
        return


def main():
    server = ThreadingHTTPServer(("0.0.0.0", PORT), Handler)
    log(f"сайт запущен:  http://localhost:{PORT}")
    log(f"Яндекс.Диск:   {YADISK_TOKEN and 'токен задан' or 'не задан (таблица остаётся локально)'}")
    log(f"n8n-вебхук:    {WEBHOOK_URL or 'не задан'}")
    log(f"заявки:        {LEADS_FILE} · таблица {LEADS_XLSX}")
    log("остановить сервер: Ctrl+C")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        log("остановлено")
        server.shutdown()


if __name__ == "__main__":
    main()